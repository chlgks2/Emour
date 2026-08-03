# -*- coding: utf-8 -*-
"""
AI Hub 감성대화 말뭉치(xlsx) → 우리 15라벨 학습용 CSV(text,label).
58개 세부감정(감정_소분류) → 15라벨 매핑. 사람문장1(첫 발화)만 사용(깨끗한 라벨).

usage:  python convert_aihub.py <input.xlsx> <output.csv>
        (경로가 상대경로면 EMOUR_DATA_DIR 기준)
"""
import os, sys, csv, collections
from openpyxl import load_workbook

DATA_DIR = os.environ.get("EMOUR_DATA_DIR", os.path.join(os.path.dirname(os.path.abspath(__file__)), "data"))
_p = lambda n: n if os.path.isabs(n) else os.path.join(DATA_DIR, n)

# 58 세부감정 → 15라벨 (실제 예문으로 검증한 매핑)
SUB_TO_LABEL = {
    "노여워하는":"화남","짜증내는":"화남","성가신":"화남","분노":"화남","툴툴대는":"화남",
    "악의적인":"화남","방어적인":"화남","구역질 나는":"화남","혐오스러운":"화남","안달하는":"걱정",
    "우울한":"슬픔","눈물이 나는":"슬픔","비통한":"슬픔","슬픔":"슬픔","후회되는":"슬픔",
    "낙담한":"슬픔","환멸을 느끼는":"슬픔","염세적인":"슬픔","외로운":"슬픔","마비된":"힘듦",
    "걱정스러운":"걱정","두려운":"걱정","초조한":"걱정","불안":"걱정","조심스러운":"걱정",
    "취약한":"걱정","회의적인":"걱정","스트레스 받는":"힘듦","혼란스러운":"당황","당혹스러운":"당황",
    "고립된":"서운함","억울한":"서운함","배신당한":"서운함","상처":"서운함","질투하는":"서운함",
    "희생된":"서운함","버려진":"서운함","실망한":"서운함","가난한, 불우한":"슬픔",
    "괴로워하는":"힘듦","좌절한":"힘듦","충격 받은":"당황",
    "당황":"당황","부끄러운":"부끄러움","남의 시선을 의식하는":"부끄러움","열등감":"부끄러움",
    "죄책감의":"미안함","한심한":"부끄러움",
    "기쁨":"기쁨","만족스러운":"기쁨","신이 난":"기쁨","자신하는":"기쁨",
    "안도":"편안","편안한":"편안","느긋":"편안","신뢰하는":"편안","감사하는":"고마움","흥분":"설렘",
}

def main():
    if len(sys.argv) < 3:
        sys.exit("usage: python convert_aihub.py <input.xlsx> <output.csv>")
    xp, out = _p(sys.argv[1]), _p(sys.argv[2])
    ws = load_workbook(xp, read_only=True).active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        sub = r[6]                     # 감정_소분류
        text = (r[7] or "").strip()    # 사람문장1
        lab = SUB_TO_LABEL.get(sub)
        if lab and len(text) >= 2:
            rows.append((text, lab))
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(["text", "label"]); w.writerows(rows)
    print(f"{len(rows)}행 → {out}")
    for l, c in collections.Counter(l for _, l in rows).most_common():
        print(f"  {l}: {c}")

if __name__ == "__main__":
    main()
